from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
VERSION_DATE = "2026 年 9 月 22 日"
FONT = "Noto Sans CJK TC"
BLUE = "24496B"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def set_cell_fill(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    # Use the legacy-compatible left/right names so the document also opens in
    # older desktop Word versions (w:start/w:end are rejected by Word 2007).
    for name, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def set_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_font(run, 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    tail = paragraph.add_run(" 頁")
    set_font(tail, 9)


def configure_document(doc: Document, short_title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18
    for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(short_title)
    set_font(run, 8.5, color=(90, 90, 90))
    add_page_number(section.footer.paragraphs[0])


def cover(doc: Document, title: str, subtitle: str, audience: str) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    run = p.add_run(title)
    set_font(run, 25, True, (0, 0, 0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_font(run, 13, False, (55, 55, 55))
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"適用對象  {audience}"), 10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"內容基準日  {VERSION_DATE}"), 10.5)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))


def add_bullets(doc: Document, items: Iterable[str], numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        set_font(p.add_run(item))


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for i, (cell, value) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = ""
        set_cell_fill(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), 9.5, True, (255, 255, 255))
        if widths:
            cell.width = Inches(widths[i])
    set_repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.text = ""
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for line_index, line in enumerate(str(value).split("\n")):
                if line_index:
                    p.add_run().add_break()
                set_font(p.add_run(line), 9.2)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph()


def section_overview(doc: Document, sections: list[tuple[str, str]]) -> None:
    add_heading(doc, "文件導覽", 1)
    add_paragraph(doc, "本頁列出主要章節與使用方式。可依目前工作直接跳到相關章節，不必依序閱讀。")
    add_table(doc, ["章節", "內容"], [[a, b] for a, b in sections], [1.75, 5.0])


def build_user_guide() -> Path:
    doc = Document()
    configure_document(doc, "Life Pilot 使用者操作說明")
    cover(doc, "Life Pilot 使用者操作說明", "完整功能 操作流程 資料安全與常見問題", "一般使用者與管理者")
    section_overview(doc, [
        ("快速開始", "註冊、登入、選擇儲存位置與首頁導覽"),
        ("個人管理", "行事曆、記帳、積分、回憶與資料匯出"),
        ("探索與學習", "推薦活動、推薦景點、地圖與遊戲題庫"),
        ("方案與額度", "免費雲端、雲端版本、此裝置版本、到期與搬移"),
        ("問題處理", "載入、同步、瀏覽器資料、額度與通知排查"),
    ])

    add_heading(doc, "1 產品與資料使用原則")
    add_paragraph(doc, "Life Pilot 將行程、收支、積分、回憶、推薦內容與學習遊戲整合在同一個帳號中。使用者可以選擇把個人資料保存在雲端或目前裝置；兩種模式一次只使用其中一種，避免同一筆資料在兩邊出現不同版本。")
    add_table(doc, ["功能", "可完成的工作", "重要提醒"], [
        ["首頁", "查看近日行程、今日與總計、推薦內容，以及快速新增記帳或積分", "各區塊可展開或收合；新增完成後首頁會刷新"],
        ["行事曆", "新增、編輯、完成、刪除、提醒與指定事件分享", "完成事件不再提醒；分享只提供唯讀"],
        ["記帳與積分", "多帳戶、分類、日期時間、小數、編輯與刪除", "刪除明細會重新計算今日與總計"],
        ["推薦與回憶", "城市篩選、地圖數量、圖片、天氣、加入行程與時間軸", "加入行事曆一律建立單日行程"],
        ["遊戲", "管理者題庫、自建題庫、答題歷程與關卡進度", "此裝置模式只能使用本機題庫"],
    ], [1.2, 3.1, 2.45])

    add_heading(doc, "2 註冊 登入與密碼重設")
    add_heading(doc, "2.1 註冊", 2)
    add_bullets(doc, [
        "在登入頁選擇註冊，輸入有效 Email 與符合畫面規則的密碼。",
        "閱讀並同意服務條款與隱私政策；完成信箱驗證後再登入。",
        "首次登入若尚未選擇儲存位置，系統預設使用雲端。儲存位置不應在未登入時顯示。",
        "同一台裝置切換不同帳號時，前一帳號的行事曆、帳戶、積分、回憶與快取都應清除後再載入。",
    ], numbered=True)
    add_heading(doc, "2.2 密碼重設", 2)
    add_bullets(doc, [
        "選擇忘記密碼並輸入註冊 Email。",
        "使用最新收到的郵件連結；Web 會直接進入重設頁，Android 與 iOS 透過 App Links 開啟。",
        "設定新密碼後重新登入。若仍回到登入頁，先關閉舊頁籤，再從最新郵件開啟。",
    ], numbered=True)

    add_heading(doc, "3 儲存位置")
    add_table(doc, ["比較項目", "雲端", "此裝置"], [
        ["可使用裝置", "登入同一帳號即可跨裝置讀取", "只在目前手機、平板、電腦或瀏覽器"],
        ["網路需求", "讀寫時需要網路", "登入與部分共用服務仍可能需要網路；個人資料讀寫以本機為主"],
        ["資料上限", "依目前有效版本額度", "有效此裝置版本期間不限筆數；到期後不可新增"],
        ["圖片", "依版本容量", "可保存在本機，不占雲端容量"],
        ["行事曆分享", "可用，依額度限制", "停用，且切換前會停止現有分享"],
        ["管理者遊戲題庫", "可依權限選用", "不可讀取，只能使用本機自建題庫"],
        ["備份風險", "由雲端帳號存取", "清除網站資料、移除 App、換裝置或換瀏覽器後不會自動出現"],
    ], [1.45, 2.7, 2.7])
    add_heading(doc, "3.1 切換流程", 2)
    add_bullets(doc, [
        "雲端轉此裝置：先下載完整個人資料，寫入本機並驗證筆數，再刪除雲端來源。",
        "此裝置轉雲端：先預檢目前方案額度；任何一項超額時整批取消，不會只上傳一部分。",
        "切換成功後，另一個儲存位置的個人資料會被清除；確認視窗會明確說明影響。",
        "Web 本機資料位於該網站的 IndexedDB。Hot Reload 不會清除；清除網站資料或改用另一瀏覽器會看不到。",
        "未來若需再次搬移，請從設定中的儲存位置功能操作，不要自行修改瀏覽器資料庫。",
    ])

    add_heading(doc, "4 方案 訂閱與額度")
    add_paragraph(doc, "方案頁最上方顯示目前儲存位置所適用的方案與實際額度。版本說明依目前方案優先排列，再依費用由低到高排列。舊版尚未到期且另購新版額度時，兩筆權益會分開列出並在有效期間加總。")
    add_table(doc, ["資源", "免費雲端", "雲端付費基準", "此裝置付費"], [
        ["行事曆", "30 筆", "300 筆 × 倍率", "不限筆數"],
        ["記帳明細", "30 筆", "300 筆 × 倍率", "不限筆數"],
        ["積分明細", "30 筆", "300 筆 × 倍率", "不限筆數"],
        ["回憶", "30 筆", "300 筆 × 倍率", "不限筆數"],
        ["自建題目", "50 題", "500 題 × 倍率", "不限題數"],
        ["行事曆分享", "2 人", "5 人 × 倍率", "不提供"],
        ["圖片", "不提供", "300 MB × 倍率", "保存在本機"],
        ["詳細答題紀錄", "近 30 天", "近 1 年", "不限保留天數"],
    ], [1.55, 1.45, 1.8, 1.8])
    add_heading(doc, "4.1 到期與超額", 2)
    add_bullets(doc, [
        "取消續訂後仍可使用到已付款期間結束。",
        "雲端付費到期後降為免費雲端；若資料仍在免費額度內，可繼續使用。",
        "超出免費額度時進入 30 天寬限期，雲端資料唯讀，可查看、匯出、搬到本機或續訂，但不可新增。",
        "寬限期結束仍未處理時，系統按目前有效額度刪除超額部分，優先保留最新資料與保留項。",
        "此裝置版本到期後，本機資料不會被系統刪除，但不能新增；續訂後恢復新增。",
        "管理者建立多筆額度時，各筆可分別刪除。只有明確按下刪除全部訂閱與額度設定，才會移除全部額度設定。",
    ])

    add_heading(doc, "5 首頁")
    add_bullets(doc, [
        "近日行程預設展開；推薦活動、推薦景點、記帳與積分預設收合。",
        "帳戶或積分選擇器用來切換帳戶；點我看更多在已選帳戶時直接進入該帳戶明細，未選時進入帳戶首頁。",
        "記帳與積分可直接新增明細，並選擇日期、時間、分類與數值。確認視窗會顯示選定日期時間。",
        "推薦活動或景點加入行事曆前會先顯示勾選與確認視窗，可調整日期時間；確認後建立單日行程。",
        "完成行程時可同時建立回憶、收入、支出、加分或減分；收支與積分共用選定的日期時間。",
        "加入行程、完成行程或新增明細後，首頁與行事曆相關月份快取會立即更新。",
    ])

    add_heading(doc, "6 行事曆與提醒")
    add_heading(doc, "6.1 事件操作", 2)
    add_bullets(doc, [
        "事件可輸入名稱、描述、國家、城市、地點、開始與結束日期時間、提醒及圖片。",
        "編輯日期時間後，手機排程通知會同步取消舊時間並建立新時間。",
        "完成事件後，Web 與手機都不應再次提醒；跨裝置完成狀態會由行事曆查詢一併取得。",
        "事件關聯記帳或積分後，行事曆與回憶走廊的快捷圖示會直接進入對應帳戶明細。",
        "月份畫面會包含第一列與最後一列跨月日期；相關月份異動時會一起清除快取。",
    ])
    add_heading(doc, "6.2 共用行事曆", 2)
    add_bullets(doc, [
        "邀請區塊、接收區塊與分享事件區塊預設收合，使用同一層外部捲動。",
        "邀請者輸入 Email，選擇全部事件或指定事件；過去與未來事件均可分享。",
        "受邀者接受後，只能查看事件，不能修改或刪除分享者資料。",
        "接受後仍可停止查看；分享者可逐筆取消事件或全部停止分享。",
        "已拒絕、已停止分享與已停止查看的紀錄不顯示，且立即釋放分享額度。",
        "重複邀請、額度已滿、已存在待接受邀請等狀態會顯示明確原因。",
    ])

    add_heading(doc, "7 記帳與積分")
    add_table(doc, ["項目", "記帳", "積分"], [
        ["一級分類", "食 衣 住 行 育 樂 保留項 未分類", "德 智 體 群 美 保留項 未分類"],
        ["二級分類", "可選填，只在編輯或匯出時顯示", "可選填，只在編輯或匯出時顯示"],
        ["數值", "收入或支出，最多四位小數並顯示千分位", "加分或減分，最多四位小數並顯示千分位"],
        ["日期時間", "新增與編輯都可修改", "新增與編輯都可修改"],
        ["刪除影響", "重算帳戶金額、今日與總計", "重算帳戶點數、今日與總計"],
    ], [1.25, 2.8, 2.8])
    add_paragraph(doc, "一般帳戶預設載入近 30 天；近 30 天沒有資料時顯示最新一筆。點選載入更多會每次再往前載入 30 天並去除重複。旅程記帳顯示全部紀錄；團體積分仍使用 30 天分批載入。保留項不受日期範圍限制。")

    add_heading(doc, "8 推薦活動 推薦景點與回憶走廊")
    add_bullets(doc, [
        "清單可依文字與城市篩選；文字符合其他城市而目前城市沒有結果時，可自動切換到符合城市。",
        "地圖模式以國家與縣市統計數量，不直接把天氣用的近似座標當成活動精確位置。點城市名稱或數字都會切回該城市清單。",
        "推薦活動先依喜歡、一般、不喜歡分組，再依日期時間、國家、城市與地點排序；不喜歡不會消失。",
        "天氣只預抓畫面可見、近期且八天內的內容；往下滑時再載入新的可見項目。",
        "只有存在子活動時才顯示預覽。從首頁或行事曆建立回憶時，只帶入當日子活動。",
        "回憶預設載入近 30 天；未來回憶也會顯示。每次載入更多再往前 30 天。",
    ])

    add_heading(doc, "9 遊戲與自建題庫")
    add_bullets(doc, [
        "雲端模式可依權限選擇管理者題庫或自己的題庫；此裝置模式只能使用本機題庫。",
        "進入遊戲前會檢查目前題庫可用題數。需要選項的遊戲以可用題目合計至少三題為原則；Social 至少要有一個啟用情境且具備三個選項。",
        "題庫管理提供搜尋、分類、啟用、停用、編輯與刪除；篩選選單變更後立即查詢。",
        "題目、描述、答案選項與回饋均支援多行輸入。",
        "遊戲首頁只載入解鎖關卡摘要、最近紀錄與是否還有更早資料，不再下載完整答題歷程。",
        "本機答題紀錄保存在目前裝置，不套用雲端 30 天或 1 年清理規則。",
    ])

    add_heading(doc, "10 個人資料匯出與帳號刪除")
    add_bullets(doc, [
        "資料匯出會下載 Excel，內容只包含行事曆、回憶走廊、記帳與積分。頁籤及欄位名稱使用當下畫面語言。",
        "匯出不會刪除任何資料。行事曆與回憶依日期時間、城市、地點與名稱排序；記帳與積分包含帳戶、分類、次分類與完整日期範圍。",
        "刪除帳號採申請制。送出後顯示申請中，使用者可撤銷；管理者確認後才真正刪除帳號與資料。",
        "儲存位置中的清理只處理個人資料或超額資料，不等同刪除登入帳號。",
    ])

    add_heading(doc, "11 常見問題")
    add_table(doc, ["狀況", "處理方式"], [
        ["遊戲頁載入較久", "確認網路與 migration 已更新。新版只查最高通關摘要；首次讀取題庫後再次進入應更快。"],
        ["本機資料看不到", "確認使用同一裝置、同一瀏覽器與同一網站網址，且未清除網站資料。"],
        ["行事曆仍顯示舊資料", "重新進入月份；若剛從首頁加入，新版會同步清除受影響月份快取。"],
        ["分享額度沒有釋放", "在分享視窗手動重新整理；拒絕、停止查看或停止分享後應立即釋放。"],
        ["額度誤建", "管理者查詢使用者後，在已建立的額度清單刪除指定一筆。不要按刪除全部訂閱與額度設定。"],
        ["無法切換到雲端", "檢查是否有有效雲端方案，以及本機資料是否超出該方案任何一項額度。"],
    ], [2.0, 4.75])

    path = OUT / "Life_Pilot_使用者操作說明.docx"
    doc.save(path)
    return path


def build_developer_guide() -> Path:
    doc = Document()
    configure_document(doc, "Life Pilot 開發者技術文件")
    cover(doc, "Life Pilot 開發者技術文件", "架構 資料模型 權限 效能 維運與發行", "開發者 維運人員與技術審查者")
    section_overview(doc, [
        ("架構", "Flutter Supabase FastAPI 與本機儲存的責任邊界"),
        ("資料", "雲端與本機模式、搬移、額度與清理"),
        ("功能", "行事曆、推薦、遊戲、記帳與積分的重要流程"),
        ("品質", "效能、生命週期、多國語系、測試與安全"),
        ("發行", "Migration、Web、Android、iOS 與 FastAPI 部署"),
    ])

    add_heading(doc, "1 系統架構")
    add_paragraph(doc, "Life Pilot 是 Flutter 多平台應用程式。Supabase 提供身分驗證、PostgreSQL、RLS 與 RPC；FastAPI 代理推薦活動、天氣、地理編碼與股票等外部服務；LocalDataStore 在 Web 使用 IndexedDB，在其他平台使用 Sembast。資料模式由 ControllerAuth 統一判斷，業務 Service 不應自行猜測資料來源。")
    add_table(doc, ["層級", "主要責任", "禁止事項"], [
        ["Widget Page", "畫面狀態、輸入、導覽、載入與錯誤呈現", "直接拼接複雜 SQL 或跨多個資料來源"],
        ["Controller", "畫面所需狀態、快取失效、非同步生命週期", "dispose 後 notifyListeners 或保留前一帳號資料"],
        ["Service", "Supabase、FastAPI、本機儲存及錯誤轉譯", "由 UI 傳入可偽造的 owner 而不核對 auth"],
        ["Supabase", "資料完整性、RLS、額度、共享與交易式操作", "把 service role 暴露給客戶端"],
        ["FastAPI", "外部網站、批次工作、service role 寫入", "信任未驗證的客戶端帳號欄位"],
    ], [1.2, 3.0, 2.55])

    add_heading(doc, "2 專案結構與開發指令")
    add_table(doc, ["路徑", "用途"], [
        ["lib/auth", "登入、註冊、Recovery、角色、訂閱快照與帳號切換清理"],
        ["lib/calendar", "事件 CRUD、月快取、提醒、分享與完成狀態"],
        ["lib/accounting  lib/point_record", "帳戶、明細、分類、日期時間、小數與總計"],
        ["lib/event  lib/memory_trace", "推薦、地圖、天氣、圖片、匯入與回憶"],
        ["lib/game", "題庫、遊戲、進度、題數可用性與本機題庫"],
        ["lib/local_storage", "IndexedDB 或 Sembast、資料模式與搬移"],
        ["lib/subscription", "方案、版本化額度、管理者工具與資料清理"],
        ["supabase/migrations", "資料表、RLS、RPC、觸發器與可稽核變更"],
        ["python_ai_service", "FastAPI、外部資料來源與後端排程"],
        ["test", "Flutter unit 與 widget tests"],
    ], [2.25, 4.5])
    add_paragraph(doc, "常用 Flutter 指令：")
    add_bullets(doc, ["flutter pub get", "flutter gen-l10n", "flutter analyze", "flutter test", "flutter run --dart-define=API_BASE_URL=http://127.0.0.1:8000"])
    add_paragraph(doc, "FastAPI 使用 requirements 安裝依賴，以 uvicorn 啟動，並用 unittest discover 執行測試。機密環境變數不得提交至 Git。")

    add_heading(doc, "3 啟動 驗證與帳號切換")
    add_bullets(doc, [
        "main 啟動先初始化 Flutter binding 與 Supabase，再顯示依賴登入狀態的畫面。SharedPreferences 等 plugin 必須在原生 plugin 註冊完成後使用。",
        "Web Recovery 只解析目前 URL；Android 與 iOS 才監聽 app_links 串流，並於 dispose 取消訂閱。",
        "帳號變更時以使用者 ID 或 Email 建立 session generation。所有帳戶、首頁、行事曆、回憶、遊戲與本機快取必須清除後再載入。",
        "同帳號 token 更新不應清除資料；只有實際帳號識別變更才重設。",
        "非同步工作完成前檢查 mounted、disposed 或 requestId，舊請求不得覆蓋新帳號或新篩選結果。",
    ])

    add_heading(doc, "4 雲端與本機資料模式")
    add_paragraph(doc, "DataStorageLocation 是個人資料唯一來源選擇。雲端與本機不可同時作為主要來源。Service 進行 CRUD 前必須查詢目前偏好；本機模式不得因 ID 空字串而向 Supabase 發出查詢。")
    add_heading(doc, "4.1 搬移交易", 2)
    add_bullets(doc, [
        "Cloud to local：RPC 匯出 → 分資源寫入 → 筆數與必要欄位驗證 → 呼叫確認刪除雲端 → 更新 preferredLocation。",
        "Local to cloud：本機清單快照 → 方案與額度預檢 → 整批 restore RPC → 雲端驗證 → 清除本機 → 更新 preferredLocation。",
        "任一資源失敗時保留來源資料，不更新 preferredLocation。NOT NULL 欄位如 country 必須在轉換時補標準預設值。",
        "Web 清理 IndexedDB 時避免混用舊版 idb_shim cursor 或跨分頁 notification 型別；以目前封裝逐資源刪除。",
    ])
    add_heading(doc, "4.2 本機額度計數", 2)
    add_paragraph(doc, "本機資料量大時不可每次開畫面逐資源完整掃描。LocalUsageCounter 以增刪異動更新快取計數；搬移、全清與帳號切換後重建。UI 顯示本機不限量，但新增權限仍依本機版本有效期判斷。")

    add_heading(doc, "5 Supabase 安全模型")
    add_bullets(doc, [
        "所有 public 個人資料表啟用 RLS。owner_id 優先使用 auth.uid()；舊表以正規化 JWT Email 比對。",
        "anon 不得執行個人資料 RPC。authenticated 僅取得必要 CRUD；REFERENCES、TRIGGER、TRUNCATE 不提供。",
        "SECURITY DEFINER 必須 SET search_path = ''，資料表與函式使用完整 schema 名稱，並在函式內再次驗證 auth.uid 或管理者角色。",
        "管理者判斷由 app_metadata.role = admin 提供，不接受使用者可修改的 user_metadata。",
        "service_role 僅存在 FastAPI 或安全後端。Flutter、Web bundle、Git、log 與錯誤訊息均不得包含。",
    ])

    add_heading(doc, "6 訂閱與版本化額度")
    add_table(doc, ["資料結構", "作用"], [
        ["subscription_pricing_versions", "管理者建立雲端或本機價格、額度、生效日與版本名稱"],
        ["user_subscriptions", "使用者目前狀態、主要版本、儲存方案、到期日與管理註記"],
        ["user_subscription_entitlements", "每次付款凍結的版本快照，可同時存在多筆並加總"],
        ["subscription_cleanup_audit", "自動或手動清理原因與各資源刪除數量"],
        ["subscription_payment_events", "付款事件與冪等處理依據"],
    ], [2.45, 4.3])
    add_bullets(doc, [
        "life_pilot_user_cloud_limit 依目前有效 entitlement 加總；無有效權益時使用免費額度。不可將額度硬寫在 Flutter。",
        "到期後雲端超額資料進入 30 天寬限；每日排程執行清理函式，只刪目前額度以上的資料。",
        "保留項在記帳與積分清理時優先保留，再以日期由新到舊保留一般明細。",
        "管理者查詢使用者後，admin_get_user_subscription_entitlements 逐筆列出有效額度。",
        "刪除誤建額度使用 admin_delete_user_subscription_entitlement(uuid)，只能依 entitlement ID 刪一筆；刪除後以剩餘有效額度重算 user_subscriptions。",
        "admin_delete_user_subscription(email) 是明確的全部訂閱設定刪除，不可綁在不具說明的垃圾桶圖示。",
    ])

    add_heading(doc, "7 行事曆 分享 快取與通知")
    add_bullets(doc, [
        "get_filtered_calendar_events 一次回傳 own 與 accepted selective share 可見事件，以及 is_completed。客戶端不需再查一次完成狀態。",
        "月快取 key 包含帳號、儲存模式與月份。新增、編輯、刪除或完成事件後，依 startDate 到 endDate 涵蓋月份失效；月格跨月列也包含在內。",
        "calendar_share_invitations 保存關係狀態，calendar_share_events 保存指定事件。拒絕、停止查看或撤銷後清理明細並釋放 active share 額度。",
        "分享視窗採手動輕量重新整理，不使用每八秒輪詢，以免使用者選取事件時清單跳動。",
        "手機通知使用事件 ID 作排程識別。日期時間或完成狀態改變時取消舊通知；App 恢復時以雲端可見事件同步。",
    ])

    add_heading(doc, "8 記帳與積分")
    add_bullets(doc, [
        "value 與 balance 使用 PostgreSQL numeric，Dart 使用 num 或 double，不可降回 bigint。輸入最多四位小數，多餘小數點或位數採字串修正而非清空整欄。",
        "日期排序以 date 為準，created_at 只作穩定次序。新增與編輯都能修改日期與時間。",
        "primary_category 是固定一級分類碼，group 是使用者二級分類。顯示與匯出時由 localization 對映，不直接輸出資料庫碼。",
        "刪除明細後在同一流程重算帳戶餘額或點數，再刷新首頁摘要。",
        "近 30 天查詢與載入更多依目前資料模式走 LocalDataStore 或 Supabase；保留項不受日期限制。",
    ])

    add_heading(doc, "9 推薦活動 景點 回憶 地圖與天氣")
    add_bullets(doc, [
        "country 使用 ISO alpha-2；map_lat/map_lng 是地圖專用座標，lat/lng 可保留天氣附近定位。已有專用座標時不重複地理編碼。",
        "推薦清單使用資料庫端分頁，篩選條件、排序與 like/dislike 在服務端一致實作。喜好更新只更新目標項目，不重新下載整批。",
        "活動去重：匯入資料無時間時，以名稱、日期、城市、地點判定；有時間時再加入開始時間。master_url 不作主要去重鍵。",
        "recommended_events_deleted 保存 tombstone，避免畫面刪除後被爬蟲重新匯入；過期清理同時處理 end_date 為空而 start_date 已過期。",
        "天氣只對可見且八天內項目載入，滑動時增量取得。HTTP client 的擁有者與關閉時機必須一致，避免 Client is already closed。",
        "回憶查詢包含未來資料，預設 30 天並分批往前；圖片延遲解碼，無圖片不保留空白區域。",
    ])

    add_heading(doc, "10 遊戲與題庫效能")
    add_bullets(doc, [
        "題庫來源由 GameQuestionBankVisibility 與 ServiceGame 統一決定。local 模式禁止 admin 題庫查詢。",
        "開始遊戲前呼叫可用性查詢。需要選項的遊戲以目前關卡可用題目合計至少三題；Social 需至少一個啟用情境且三個 choices。",
        "遊戲清單先載入 game_list，再平行查最高通關摘要、近 30 天紀錄與更早資料存在性。",
        "get_user_game_highest_passed_level 只回傳 max(level)，不再透過 fetch_user_progress 下載完整歷程。",
        "本機 _localProgressRequests 讓同一輪的摘要、列表與 hasMore 共用一次 IndexedDB 掃描；每次重新載入前失效，避免看不到新成績。",
        "Migration 尚未部署時針對 42883 或 PGRST202 暫時 fallback 舊查詢，部署完成後走摘要 RPC。",
    ])

    add_heading(doc, "11 多國語系 UI 與生命週期")
    add_bullets(doc, [
        "所有使用者可見字串放在 zh、en、ja、ko ARB，修改後執行 flutter gen-l10n。資料庫分類碼透過 loc 對映。",
        "手機與 Web 使用 LayoutBuilder、Wrap、Expanded 與約束寬度避免 RenderFlex overflow。標題必須完整顯示，控制項不足時換行。",
        "清單使用 builder、itemExtent 或 prototypeItem、cacheExtent 與圖片尺寸快取。多圖片對話框只在可見時解碼。",
        "任何 timer、stream、HTTP、Future 回呼都必須在 dispose 後停止；setState callback 不可標記 async。",
        "文字 formatter 回傳的新 TextEditingValue 必須讓 selection 和 composing range 位於新字串長度內。",
    ])

    add_heading(doc, "12 Migration 與部署")
    add_bullets(doc, [
        "Migration 檔名採 UTC 時序前綴，使用 transaction。變更 RETURNS TABLE 時先以完整 signature DROP 再 CREATE。",
        "部署順序：資料庫 migration → RPC 權限查核 → FastAPI → Flutter Web 或 App。客戶端先上線而 RPC 未部署時，只允許有明確 fallback 的功能。",
        "Web CI 執行 pub get、gen-l10n、analyze、test 與 build web。pub.dev authorization failure 應先檢查 CI token 或移除錯誤的私有 token。",
        "Android Gradle wrapper 必須符合目前 Flutter 最低版本；bundle ID 與 iOS identifier 均為 com.minavi.life_pilot。",
        "Android release key、key.properties、Apple signing certificate 與 service role 均由安全環境管理。",
    ])

    add_heading(doc, "13 最新 Migration")
    add_table(doc, ["檔案", "內容", "App 依賴"], [
        ["20260922100000_add_game_progress_summary.sql", "新增最高通關關卡摘要 RPC 與 authenticated 權限", "遊戲頁載入效能"],
        ["20260922110000_admin_manage_individual_entitlements.sql", "新增管理者額度清單與依 ID 刪除單筆額度", "管理使用者訂閱"],
    ], [2.7, 2.85, 1.2])

    add_heading(doc, "14 維運檢查")
    add_bullets(doc, [
        "每日檢查推薦活動更新摘要、失敗來源、候選數與寫入數。",
        "每日排程 subscription overage cleanup 與 game answer history cleanup，確認 service_role only。",
        "定期稽核 anon grants、SECURITY DEFINER search_path、RLS policy 與 orphan records。",
        "監控 Supabase 資料量、圖片容量、RPC 延遲、FastAPI 5xx、Web 首屏與遊戲頁載入時間。",
        "發布前執行完整測試文件的 Gate，不以 analyze 通過取代跨裝置手動驗收。",
    ])

    path = OUT / "Life_Pilot_開發者技術文件.docx"
    doc.save(path)
    return path


def add_test_case(doc: Document, case_id: str, title: str, pre: str, steps: list[str], expected: list[str], checks: list[str]) -> None:
    add_heading(doc, f"{case_id} {title}", 2)
    rows = [
        ["前置條件", pre],
        ["操作步驟", "\n".join(f"{i + 1}. {step}" for i, step in enumerate(steps))],
        ["預期結果", "\n".join(f"• {item}" for item in expected)],
        ["資料檢查點", "\n".join(f"• {item}" for item in checks)],
        ["實際結果", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"],
        ["測試結果", "□ 通過    □ 失敗    □ 阻塞    缺陷編號 ＿＿＿＿＿＿"],
    ]
    add_table(doc, ["欄位", "內容"], rows, [1.35, 5.4])


def build_test_plan() -> Path:
    doc = Document()
    configure_document(doc, "Life Pilot 測試計畫與檢查點")
    cover(doc, "Life Pilot 測試計畫與檢查點", "功能流程 資料驗證 權限 效能與上線驗收", "測試人員 開發者與產品負責人")
    section_overview(doc, [
        ("測試準備", "環境、帳號、資料與自動化基線"),
        ("核心流程", "登入、儲存位置、方案、行事曆、記帳與積分"),
        ("內容功能", "推薦、回憶、遊戲、匯出與帳號刪除"),
        ("非功能", "效能、生命週期、多國語系、響應式與安全"),
        ("上線 Gate", "各平台發布前必須全部完成的檢查"),
    ])

    add_heading(doc, "1 測試執行方式")
    add_paragraph(doc, "每個案例都要記錄實際結果、通過狀態與缺陷編號。涉及雲端與本機搬移、刪除、額度或分享的案例，除了畫面結果，必須同時檢查資料庫或 IndexedDB。破壞性案例只使用測試帳號。")
    add_table(doc, ["維度", "至少涵蓋"], [
        ["平台", "Chrome 寬螢幕與窄螢幕、Android 手機、Android 平板、iPhone、iPad"],
        ["語系", "繁體中文、英文、日文、韓文"],
        ["帳號", "免費雲端、雲端付費、此裝置付費、管理者、到期、超額、待刪除"],
        ["網路", "正常、慢速、離線、API timeout、Supabase 4xx 或 5xx"],
        ["資料量", "空資料、額度邊界、大量清單、跨月、含圖片與長文字"],
    ], [1.45, 5.3])
    add_paragraph(doc, "自動化基線：flutter pub get、flutter gen-l10n、flutter analyze、flutter test；FastAPI 執行 compileall 與 unittest。")

    cases = [
        ("AUTH-01", "新帳號註冊與首次登入", "準備未註冊 Email，使用雲端預設方案。", ["完成註冊與信箱驗證", "登入後進入首頁", "登出再登入"], ["首次登入不顯示前一帳號資料", "未設定儲存位置時使用雲端", "首頁載入完成且無無限轉圈"], ["auth.users 僅一筆", "Controller 快取 owner 為新帳號", "dashboard_setting 不引用前一帳號"]),
        ("AUTH-02", "Web 密碼重設", "準備可收信帳號與部署網址。", ["申請重設密碼", "點最新郵件連結", "輸入新密碼並送出", "以新密碼登入"], ["連結進入重設頁而非一般登入頁", "無 main.dart.js 未捕捉錯誤", "新密碼可登入"], ["Recovery URL fragment 解析成功", "Web 未啟動原生 app_links listener"]),
        ("AUTH-03", "跨帳號資料隔離", "帳號 A 有行事曆、記帳、積分與回憶；帳號 B 為新帳號。", ["登入 A 並瀏覽各頁", "登出 A", "登入 B，立即開首頁與行事曆"], ["B 看不到 A 的事件、帳戶、積分或快取", "不需要手動換月份才恢復"], ["session generation 已更新", "各 Controller 已清空", "Supabase request 使用 B token"]),
        ("STORAGE-01", "雲端轉此裝置", "雲端帳號四類個人資料均有資料，且具有效此裝置方案。", ["在儲存位置選此裝置", "閱讀差異並確認", "等待搬移完成", "重開各功能頁"], ["所有支援資源顯示相同筆數", "雲端來源在驗證後刪除", "設定頁自動關閉"], ["IndexedDB 或 Sembast 筆數", "雲端個人表為零", "preferredLocation=local"]),
        ("STORAGE-02", "此裝置轉雲端額度內", "本機資料均小於目前有效雲端額度。", ["選擇雲端", "確認將刪除本機資料", "完成後重開 App"], ["資料完整上傳", "本機個人資料清除", "方案仍為原有效雲端方案"], ["雲端各表筆數", "本機 owner 資源為零", "entitlement 未被改為免費"]),
        ("STORAGE-03", "此裝置轉雲端超額", "其中一項本機資料超過雲端額度。", ["選擇雲端", "查看預檢結果", "取消或確認"], ["整批上傳不開始", "顯示超額資源、已用、額度與差額", "本機資料完整保留"], ["雲端筆數未增加", "preferredLocation 仍為 local"]),
        ("SUB-01", "免費雲端額度邊界", "準備免費帳號，逐項建立至上限前一筆。", ["建立最後一筆額度內資料", "再新增一筆", "刪除一筆後重新新增"], ["上限內成功", "超額顯示明確可用筆數與方案提示", "刪除後立即釋放額度"], ["get_my_subscription_usage used/quota", "實際表筆數"]),
        ("SUB-02", "多筆付費額度加總", "同一使用者具有舊版未到期額度與新版加購額度。", ["開啟方案頁", "檢查每筆版本", "新增至合計額度邊界"], ["目前方案优先顯示", "兩筆版本、費用、生效日與到期日分開顯示", "使用上限為有效 entitlement 加總"], ["user_subscription_entitlements 兩筆", "life_pilot_user_cloud_limit 等於快照加總"]),
        ("SUB-03", "管理者刪除單筆誤建額度", "替同一使用者建立至少三筆有效額度。", ["管理者查詢使用者", "在已建立的額度清單刪除中間一筆", "重新查詢並登入使用者帳號"], ["只刪指定版本", "其餘兩筆保留", "目前方案依剩餘有效額度重算"], ["deleted entitlement ID 不存在", "其他 entitlement ID 仍存在", "user_subscriptions 到期日為剩餘最晚到期日"]),
        ("SUB-04", "管理者刪除全部訂閱設定", "測試使用者有多筆額度，但個人資料不可刪除。", ["按紅色完整文字按鈕", "確認警告", "重新查詢"], ["全部 entitlement 與 subscription 設定移除", "行事曆、記帳、積分與回憶仍存在"], ["訂閱表為零", "個人資料表筆數不變"]),
        ("SUB-05", "到期與寬限期", "雲端付費帳號超出免費額度並模擬到期。", ["刷新訂閱狀態", "嘗試新增", "查看超額項目與截止日"], ["資料唯讀", "顯示寬限截止日與超額明細", "可匯出、搬移或續訂"], ["downgrade_grace_ends_at", "新增 RPC 被拒絕", "清理前資料仍完整"]),
        ("CAL-01", "事件 CRUD 與跨月快取", "選擇月底日期並準備跨月事件。", ["新增跨月事件", "切換前後月份", "修改到另一月份", "刪除事件"], ["涉及月份都即時更新", "首頁近日行程同步", "刪除後不殘留"], ["月份 cache keys 已失效", "calendar_events 日期正確"]),
        ("CAL-02", "完成事件與跨裝置提醒", "Web 與手機登入同一雲端帳號，事件有今日提醒。", ["在 Web 完成事件", "手機恢復 App", "等待原提醒時間"], ["手機同步完成狀態", "原生通知取消", "Web 提醒清單不再顯示"], ["RPC is_completed=true", "通知排程無該 event ID"]),
        ("CAL-03", "指定事件分享與唯讀", "A 有過去與未來事件，B 為一般帳號。", ["A 邀請 B 並選指定事件", "B 接受", "B 查看行事曆", "B 嘗試修改"], ["只顯示指定事件", "以顏色區分 owner", "B 無法修改"], ["invitation=accepted", "calendar_share_events ID 正確", "RLS update 拒絕"]),
        ("CAL-04", "分享拒絕 重邀與額度釋放", "A 邀請 B，B 尚未處理。", ["B 拒絕", "A 手動刷新分享視窗", "A 再次邀請 B", "B 停止查看"], ["拒絕紀錄不顯示", "額度釋放", "新邀請顯示待接受而非拒絕", "停止查看後再次釋放"], ["active status 只計 pending/accepted", "舊明細已清理"]),
        ("RECORD-01", "記帳日期時間分類與小數", "準備可用記帳帳戶。", ["新增 1,234.5678 並選日期時間與分類", "編輯時間與次分類", "刪除明細"], ["最多四位小數", "多一個小數點只移除多餘字元", "刪除後今日與總計重算"], ["numeric 值正確", "date 含選定時間", "account balance 重算"]),
        ("RECORD-02", "積分加減分與保留項", "準備團體積分帳戶與超過 30 天資料。", ["新增加分與減分", "建立保留項", "載入近 30 天與更多"], ["文案使用加分與減分", "保留項始終顯示", "載入更多不重複"], ["primary_category=reserved", "points 重算正確"]),
        ("HOME-01", "首頁帳戶切換與快捷入口", "建立多個個人帳戶與旅程或團體帳戶。", ["點選帳戶選擇器", "切換帳戶", "點點我看更多", "新增明細後返回"], ["選擇器可換帳戶", "有選帳戶時直達對應明細", "首頁摘要立即更新"], ["dashboard setting owner 正確", "未發出 id=eq. 空值請求"]),
        ("EVENT-01", "推薦篩選 地圖與加入行程", "準備多城市活動與景點。", ["輸入關鍵字", "切換地圖並點城市名稱與數字", "調整日期時間後加入行事曆"], ["沒有目前城市結果時切到符合城市", "兩種點擊都回到清單", "建立單日行程且首頁刷新"], ["country/city 正確", "calendar end_date 為空或同日規則一致", "月份 cache 已失效"]),
        ("EVENT-02", "推薦活動排序", "準備 like、一般、dislike，包含已開始及未開始事件。", ["開啟清單", "變更 like/dislike", "清除篩選"], ["like 最前、dislike 最後", "start date 早於今天視同今天", "再依 start time、end date、國家、城市、地點排序"], ["喜好更新未重新下載整批", "無 end date 時以 start date 排序"]),
        ("EVENT-03", "天氣與圖片延遲載入", "準備八天內與八天後事件，部分有圖片。", ["進入清單", "往下滑", "開啟多圖片預覽"], ["只抓可見且八天內天氣", "後續可見項目增量載入", "無圖片不占位置且圖片不造成長時間卡頓"], ["天氣 request 數量", "圖片解碼尺寸", "離頁後無 dispose callback"]),
        ("MEMORY-01", "回憶分批載入與預覽", "準備過去 70 天、今日與未來回憶，部分含 subEvents。", ["進入回憶走廊", "載入更多兩次", "開啟有 subEvents 的預覽"], ["初始近 30 天且未來項目可見", "每次往前 30 天", "預覽正常且沒有 ProviderNotFound"], ["資料去重", "本機與雲端分支一致"]),
        ("GAME-01", "遊戲頁初始載入效能", "準備大量 game_user 歷程與多關卡。", ["清除 App 畫面快取", "進入遊戲頁", "切換遊戲", "載入更多歷程"], ["初始只查最高通關摘要與近 30 天", "畫面不因完整歷程下載長時間空白", "解鎖關卡正確"], ["呼叫 get_user_game_highest_passed_level", "未呼叫完整 fetch_user_progress 作解鎖", "RPC 延遲紀錄"]),
        ("GAME-02", "本機題庫可用性", "此裝置模式，分別準備 0、2、3 題可用題目。", ["各狀態嘗試進入遊戲", "停用其中一題", "重新啟用"], ["未達最低題數顯示明確訊息", "合計至少三題時可進入", "不查管理者題庫"], ["IndexedDB 題數", "無管理者題庫 Supabase request"]),
        ("GAME-03", "本機進度共用掃描", "本機有大量答題歷程。", ["進入遊戲頁", "切換遊戲後返回", "完成一局再返回"], ["同一輪摘要、清單、hasMore 共用掃描", "新成績可見，不受舊快取影響"], ["IndexedDB list 呼叫次數", "invalidateLocalProgress 已執行"]),
        ("EXPORT-01", "多語系 Excel 匯出", "四類資料均有跨日期與分類內容。", ["依四種語系分別匯出", "開啟 Excel", "核對欄位與排序"], ["只有行事曆、回憶、記帳、積分頁籤", "頁籤與標題依語系", "日期不受 30 天限制，分類與次分類完整"], ["沒有預設 Sheet1", "金額與點數有千分位", "記帳含幣別"]),
        ("DELETE-01", "帳號刪除申請 撤銷與管理者完成", "一般雲端帳號與管理者帳號。", ["使用者送出申請", "再次申請", "使用者撤銷", "管理者確認撤銷", "重新申請並由管理者刪除"], ["首次顯示申請中而非已刪除", "再次申請顯示申請中", "撤銷後可重新申請", "管理者完成顯示帳號與相關資料已刪除"], ["申請與撤銷紀錄狀態", "auth user 及個人資料刪除結果"]),
        ("UI-01", "手機與 Web 響應式", "使用窄手機、平板與 1366px Web。", ["逐頁開啟並展開所有區塊", "輸入長標題與長帳戶名稱", "開啟地圖與分享視窗"], ["無 right 或 bottom overflow", "標題完整顯示", "控制項不足時換行，手機捲軸不蓋文字"], ["console 無 RenderFlex overflow", "Dialog 可捲動且可按確認"]),
        ("LIFE-01", "離頁後非同步工作", "使用慢速網路開啟股票、推薦、回憶與行事曆。", ["開始載入後立即離頁", "重複進出", "登出"], ["無 used after disposed", "無 setState after dispose", "無 Client already closed", "舊帳號結果不覆蓋新帳號"], ["timer/stream 已取消", "requestId 或 mounted guard 生效"]),
        ("SEC-01", "RLS 與 RPC 權限", "準備 anon、一般 authenticated、管理者與 service role。", ["逐角色查詢與寫入主要表", "嘗試執行管理 RPC", "檢查函式設定"], ["anon 無個人資料權限", "一般使用者只能自己的資料與接受分享的唯讀事件", "管理 RPC 僅管理者成功"], ["RLS enabled", "search_path 固定", "grants 無 REFERENCES/TRIGGER/TRUNCATE"]),
    ]
    for case in cases:
        add_test_case(doc, *case)

    add_heading(doc, "3 上線前 Gate")
    add_bullets(doc, [
        "所有自動化測試通過，且沒有跳過失敗案例。",
        "兩個 20260922 migration 已執行並驗證 authenticated 與 anon 權限。",
        "Web、Android 與 iOS 至少各完成一次註冊、登入、重設密碼、首頁、行事曆及資料新增 smoke test。",
        "雲端與此裝置雙向搬移、額度內、超額與中途失敗均通過。",
        "管理者逐筆刪除額度只刪指定 entitlement，其他額度與個人資料保留。",
        "遊戲頁在大量歷程下的首屏時間符合產品可接受範圍，且解鎖關卡正確。",
        "四語系無未翻譯 key、簡體中文、截字或 overflow。",
        "正式環境不含 service role、資料庫密碼、簽章密碼或測試帳號。",
    ])

    path = OUT / "Life_Pilot_測試計畫與檢查點.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [build_user_guide(), build_developer_guide(), build_test_plan()]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
